"""把 scripts/export-dataset.ts 导出的 JSON 排成 Excel：指标信息集、异动记录、Agent 评测事件。
用法：npx tsx scripts/export-dataset.ts > dataset.json && python scripts/build_dataset_xlsx.py dataset.json docs/异动监测信息集与评测事件.xlsx   （依赖 openpyxl）"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

d = json.load(open(sys.argv[1], encoding='utf-8'))
AS_OF = d['asOf']
INK, MUTED, LINE = 'FF1F2328', 'FF657689', 'FFE3E6EA'
FILL = {'head': 'FFF3F4F6', 'hit': 'FFE3F4E8', 'miss': 'FFFBE4E4', 'void': 'FFEFEFEF', 'pending': 'FFEDE9FB', 'live': 'FFE3F4E8', 'event': 'FFE1EEFB'}
STATUS = {'hit': '说对', 'miss': '说错', 'void': '作废', 'pending': '在等结果'}
KIND = {'rule': '已接入 · 日线计算', 'event': '已接入 · 人工整理事件', 'planned': '未接入'}
MARKET = {'A': 'A股', 'HK': '港股', 'US': '美股'}
font = lambda **k: Font(name='Arial', size=k.pop('size', 10), color=k.pop('color', INK), **k)
thin = Side(style='thin', color=LINE)
md = lambda s: f"{int(s[5:7])}月{int(s[8:10])}日" if s else ''
sg = lambda v: '' if v is None else f"{'+' if v > 0 else ''}{v:.2f}%"

wb = Workbook()
wb.remove(wb.active)

def sheet(name, title, sub, head, rows, widths, freeze='C6', fills=None):
    """和原表同一版式：第 2 行标题、第 3 行说明、第 5 行表头，数据从第 6 行起"""
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws['A2'] = title; ws['A2'].font = font(size=16, bold=True)
    ws['A3'] = sub; ws['A3'].font = font(color=MUTED)
    for c, h in enumerate(head, 1):
        x = ws.cell(5, c, h); x.font = font(color=MUTED, bold=True); x.fill = PatternFill('solid', fgColor=FILL['head'])
        x.alignment = Alignment(vertical='center', wrap_text=True); x.border = Border(bottom=thin, top=thin)
    ws.row_dimensions[5].height = 30
    for r, row in enumerate(rows, 6):
        for c, v in enumerate(row, 1):
            x = ws.cell(r, c, v); x.font = font(bold=(c <= 2 and isinstance(v, str) and len(v) < 24))
            x.alignment = Alignment(vertical='top', wrap_text=True, horizontal='right' if isinstance(v, (int, float)) else 'left'); x.border = Border(bottom=thin)
        if fills and fills[r - 6]:
            col, key = fills[r - 6]
            ws.cell(r, col).fill = PatternFill('solid', fgColor=FILL[key])
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A5:{get_column_letter(len(head))}{5 + len(rows)}"
    return ws

inds, anos, jus = d['indicators'], d['anomalies'], d['judgments']
live = [i for i in inds if i['kind'] != 'planned']
settled = [j for j in jus if j['status'] in ('hit', 'miss')]
hit = sum(j['status'] == 'hit' for j in jus)

# ── 说明 ──
ws = wb.create_sheet('说明'); ws.sheet_view.showGridLines = False
ws.column_dimensions['A'].width = 22; ws.column_dimensions['B'].width = 120
ws['A2'] = '异动监测信息集与 Agent 评测事件'; ws['A2'].font = font(size=16, bold=True)
ws['A3'] = f'由产品原型「异动管家」导出 · 数据截至 {md(AS_OF)}（2025 年）· 8 只自选股（A股 3、港股 2、美股 3）'; ws['A3'].font = font(color=MUTED)
notes = [
    ('这份表是什么', '在《异动监测指标设计与分级规则》的 38 项指标基础上，补上三样东西：每项指标在产品里的接入状态和回测结果；按这些指标在真实日线上实际捕捉到的异动记录；管家发出的每条判断从捕捉到复盘的全过程和对答案结果。'),
    ('指标信息集', f'{len(inds)} 项指标，已接入 {len(live)} 项。三档提醒门槛、固定等级、影响、例子沿用原表；新增大白话解释、接入状态和平衡型门槛下的回测。'),
    ('异动记录', f'{len(anos)} 条。行情指标类 {sum(a["cat"] == "价量" for a in anos)} 条和跨市场传导 {sum(a["cat"] == "跨市场传导" for a in anos)} 条由规则在日线上自动识别（3月20日起）；宏观政策与公司事件 {sum(not a["auto"] for a in anos)} 条为人工整理。'),
    ('Agent 评测事件', f'{len(jus)} 条判断，每条 5 个步骤（捕捉异动 → 传导路径 → 给出可验证的判断 → 到期对答案 → 复盘与经验），共 {len(jus) * 5} 步。已有结果的 {len(settled)} 条里说对 {hit} 条、说错 {len(settled) - hit} 条；另有作废 {sum(j["status"] == "void" for j in jus)} 条、在等结果 {sum(j["status"] == "pending" for j in jus)} 条。'),
    ('和原表的区别', '原表的评测事件是 5 条合成的参考 Trace，未运行。这一版用 2025 年 3–4 月的真实事件和真实收盘价，「对答案」一步是实际算出来的。'),
    ('哪些是真实的', '日线行情（东方财富，前复权）、指标触发、判断的基准价与结算结果，都由规则在真实数据上计算。'),
    ('哪些不是', '事件的标题、传导路径、证据摘要和复盘文字是人工整理的，没有调用大模型生成，来源也没有逐条核对；指标阈值沿用原表的候选值，没有校准；判断的幅度和期限是看过数据之后选的，不能当作样本外的准确率。'),
    ('样本量', '8 只股票、不到 4 个月，样本很小，所有比例只能说明这套流程能跑通，不能说明指标有效。'),
    ('怎么重新生成', 'npx tsx scripts/export-dataset.ts > dataset.json，再 python scripts/build_dataset_xlsx.py dataset.json <输出路径>'),
]
for r, (k, v) in enumerate(notes, 5):
    ws.cell(r, 1, k).font = font(bold=True); ws.cell(r, 2, v).font = font()
    for c in (1, 2): ws.cell(r, c).alignment = Alignment(vertical='top', wrap_text=True); ws.cell(r, c).border = Border(bottom=thin)

# ── 指标信息集 ──
def ev(i):
    e = i['eval']
    if not e: return ['', '', '', '']
    if e['kind'] == 'rule': return [e['n'], e['cont'], f"{round(e['cont'] / e['n'] * 100)}%" if e['n'] else '', sg(e['avg5']) if e['n'] else '']
    return [e['n'], f"说对 {e['hit']} / 说错 {e['miss']} / 作废 {e['void']} / 在等 {e['pending']}", '', '']
rows = [[i['id'], i['scope'], i['spec']['layer'], i['spec']['name'], i['intro'], KIND[i['kind']], i['need'], i['spec']['steady'], i['spec']['balanced'], i['spec']['aggressive'], i['spec']['levels'], i['spec']['impact'], *ev(i), i['spec']['example'], i['spec']['exclude']] for i in inds]
sheet('指标信息集', '异动监测指标 · 内容、门槛与回测', f'{len(inds)} 项 | 已接入 {len(live)} 项 | H—J 为三档提醒门槛（用户可逐项切换、开关），K 为固定等级 | 回测按平衡型门槛，截至 {md(AS_OF)}，阈值未校准',
      ['ID', '宏观/微观', '指标层级', '指标', '大白话解释', '接入状态', '缺什么数据', '稳健型 · 更早提醒', '平衡型 · 默认', '进取型 · 容忍波动', '异动等级 · 不随档位变', '可能影响 · 非涨跌预测', '回测触发次数 / 判断条数', '5日内同向再走5%以上（次）/ 判断结果', '延续比例', '触发后5日平均涨跌', '例子（合成）', '反例与误报排除'],
      rows, [7, 11, 16, 24, 52, 20, 26, 32, 32, 32, 40, 36, 12, 22, 10, 12, 46, 42], 'E6',
      [(6, 'live' if i['kind'] == 'rule' else 'event') if i['kind'] != 'planned' else None for i in inds])

# ── 八类政策 ──
sheet('八类政策', 'M03 产业监管与贸易政策 · 八类', '沿用原表。四门槛：正式证据 → 实质变化 → 公司实际适用 → 重要影响。',
      ['ID', '政策类别', '引入的内容', '影响传导', '触发判断', '合成例子'], [[p['id'], p['name'], p['content'], p['path'], p['trigger'], p['example']] for p in d['policies']], [10, 24, 52, 40, 50, 56])

# ── 异动记录 ──
rows = []
for n, a in enumerate(anos, 1):
    mv = a['moves']
    rows.append([f"A{n:02d}", a['date'], a['scope'], a['cat'], a['signal'], ' / '.join(sorted({h['level'] for h in a['indicators']})) or '—', a['strength'], '风险' if a['tone'] == 'risk' else '利好或缓解',
                 a['title'], '；'.join(f"{h['id']} {h['label']}" for h in a['indicators']) or '；'.join(f"{p['via']} → {'、'.join(p['names'])}" for p in a['paths']),
                 '、'.join(f"{m['name']}（{MARKET[m['market']]}）" for m in mv), '；'.join(f"{m['name']} {sg(m['p'])}" for m in mv), '；'.join(f"{m['name']} {sg(m['fwd5'])}" for m in mv if m['fwd5'] is not None),
                 '；'.join(f"{e['text']}〔{e['source']}〕" for e in a['evidence']), '自动 · 规则识别' if a['auto'] else '人工整理', '、'.join(a['judgments']) or '—'])
sheet('异动记录', '异动记录 · 这套指标实际捕捉到了什么', f'{len(anos)} 条，按日期排列 | 涨跌为事件可见后第一个交易日的涨跌，「之后5日」为再往后 5 个交易日的累计涨跌 | 截至 {md(AS_OF)}',
      ['编号', '日期', '宏观/微观', '类别', '主指标', '等级', '强度(0-100)', '性质', '异动', '命中指标 / 传导路径', '涉及自选股', '当日涨跌', '之后5日累计涨跌', '依据与来源', '识别方式', '对应判断'],
      rows, [7, 12, 10, 12, 8, 9, 10, 11, 46, 56, 26, 28, 28, 60, 16, 16], 'C6')

# ── 评测事件（一条判断一行）──
order = {j['id']: f"T{n:02d}" for n, j in enumerate(jus, 1)}
def outcome(j):
    far = f"窗口内收盘最{'高' if j['up'] else '低'} {sg(j['worst'])}" if j['worst'] is not None else '窗口内还没有收盘数据'
    if j['status'] == 'hit': return f"{md(j['settledOn'])} 收盘{'涨过' if j['up'] else '跌破'} {j['line']}，第 {j['elapsed']} 个交易日。{far}。"
    if j['status'] == 'miss': return f"{j['n']} 个交易日走完，没有{'涨过' if j['up'] else '跌破'} {j['line']}。{far}。"
    if j['status'] == 'void': return f"{md(j['settledOn'])} 失效条件触发：{j['voidOn']['text'] if j['voidOn'] else ''}。不计入战绩。"
    return f"已过 {j['elapsed']} / {j['n']} 个交易日，{far}；{md(j['dueDate']) + ' 到期' if j['dueDate'] else '到期日超出数据范围'}。"
rows = [[order[j['id']], j['issued'], j['event'], j['signal'], f"{j['name']}（{MARKET[j['market']]}）", '看涨' if j['up'] else '看跌', j['text'], j['invalid'], outcome(j), STATUS[j['status']],
         (j['review'] or {}).get('why', ''), (j['review'] or {}).get('missed', ''), (j['review'] or {}).get('lesson', '')] for j in jus]
sheet('评测事件', 'Agent 评测事件 · 每条判断的输入、输出与结果', f'{len(jus)} 条判断 | 已有结果 {len(settled)} 条：说对 {hit}、说错 {len(settled) - hit} | 结果由真实收盘价结算，截至 {md(AS_OF)} | 步骤明细见「评测步骤」',
      ['Trace ID', '发出日', '事件', '指标', '判断对象', '方向', '管家的判断（可验证）', '事前写明的失效条件', '实际结果', '判定', '为什么对 / 错', '当时漏看了什么', '记下的经验'],
      rows, [9, 12, 42, 8, 18, 8, 52, 40, 52, 10, 46, 46, 46], 'D6', [(10, j['status']) for j in jus])

# ── 评测步骤（一条判断五步）──
rows, fills = [], []
for j in jus:
    t, auto = order[j['id']], j['anomaly'].startswith(('auto-', 'cross-'))
    chain = '；'.join(f"{c['label']}：{c['text']}" for c in j['chain'])
    evd = '；'.join(f"{e['text']}〔{e['source']}〕" for e in j['evidence'])
    done = j['status'] != 'pending'
    steps = [
        ('捕捉异动', f"{'日线行情（截至发出日前一个交易日）' if auto else '政策与公司公告原文'}：{j['when']}", f"归入 {j['signal']}。{evd}", '归类正确；依据是正式来源或真实行情；没有用到发出日之后的信息', '已运行 · 规则计算' if auto else '人工整理 · 来源未逐条核对'),
        ('传导路径', f"自选清单 + {chain}", f"{'；'.join(j['paths']) or chain} → {j['name']}", '路径每一环都能说清；只落在自选股上；影响大小未知时不编数字', '已运行 · 规则计算' if auto else '人工整理'),
        ('给出可验证的判断', f"基准价：{md(j['baseDate'])} 收盘 {j['base']}；幅度 {j['thr']:+g}%，期限 {j['n']} 个交易日", f"{j['text']}。失效条件：{j['invalid']}", '有明确价格、期限和失效条件；基准价早于发出日', '已运行 · 幅度与期限为人工设定'),
        ('到期对答案', f"{j['name']} {md(j['issued'])} 起 {j['n']} 个交易日的真实收盘价", outcome(j), f"判定：{STATUS[j['status']]}", '已运行 · 真实收盘价结算' if done else '未到期'),
        ('复盘与经验', '判断、实际走势、同期其他事件', f"原因：{j['review']['why']} 漏看：{j['review']['missed']} 经验：{j['review']['lesson']}" if j['review'] else '还没有结果，暂不复盘。', '说对说错都要有原因和一条经验；经验由用户决定记不记进长期记忆', '人工整理 · 未调用模型' if j['review'] else '未运行'),
    ]
    for n, (stage, inp, out, chk, st) in enumerate(steps, 1):
        rows.append([t, j['event'] if n == 1 else '', f"{t}-S{n}", stage, inp, out, chk, st]); fills.append((7, j['status']) if n == 4 else None)
sheet('评测步骤', 'Agent 评测步骤 · 输入、输出与检查点', f'{len(jus)} 条判断 × 5 步 = {len(rows)} 步 | 版式沿用原表「评测事件」| 运行状态如实标注：哪些是规则算的，哪些是人工整理的',
      ['Trace ID', '事件', '步骤ID', '步骤 / 阶段', '输入', '管家输出 / 实际结果', '评测检查点', '运行状态'], rows, [9, 34, 11, 20, 60, 80, 54, 28], 'E6', fills)

wb.move_sheet('说明', -(len(wb.sheetnames) - 1))
wb.save(sys.argv[2])
print('已写入', sys.argv[2], {ws.title: ws.max_row - 5 for ws in wb if ws.title != '说明'})
